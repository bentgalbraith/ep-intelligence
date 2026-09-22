"""Compare an estate-planning PowerPoint diagram against Word drafts."""

import io
import logging
import os
import re
import time
import traceback
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

log = logging.getLogger("compare_diagram_drafts")

MAX_WORD_DOCS = 20
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TOTAL_CHARS = 250_000
MAX_ZIP_MEMBERS = 4000
MAX_UNCOMPRESSED_BYTES = 80 * 1024 * 1024

TOOL_KEY = "compare_diagram_drafts"

COMPARE_PROMPT = """\
You are a legal assistant specializing in estate planning.{firm_context}

An attorney uploaded an estate-planning diagram (PowerPoint) and one or more \
draft Word documents. Compare the drafts to the diagram.

Decide whether the drafts match the structure and the level of detail laid \
out in the PowerPoint — parties, shares, trusts, gifts, tax elections, \
fiduciaries, and other plan elements shown on the slides.

Write one paragraph for the attorney. Be specific: say what matches, what is \
missing or inconsistent, and what appears in the drafts but not on the \
diagram. If the extracted text is unclear or incomplete, say so. Do not \
invent facts. Do not give legal advice.

Return only the paragraph — no title, heading, or bullet list.
"""


class CompareError(Exception):
    """User-facing problem with the uploaded files or extracted text."""


def is_pptx_bytes(data):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = zf.namelist()
    except zipfile.BadZipFile:
        return False
    return "ppt/presentation.xml" in names or any(n.startswith("ppt/") for n in names)


def is_docx_bytes(data):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = zf.namelist()
    except zipfile.BadZipFile:
        return False
    return "word/document.xml" in names


def zip_is_oversized(data):
    """True when an Office zip would expand past our in-memory limits."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
    except zipfile.BadZipFile:
        return False
    if len(infos) > MAX_ZIP_MEMBERS:
        return True
    total = 0
    for info in infos:
        total += max(info.file_size, 0)
        if total > MAX_UNCOMPRESSED_BYTES:
            return True
    return False


def safe_filename(name):
    cleaned = re.sub(r"[\r\n\t]+", " ", name or "").strip()
    return (cleaned or "Document")[:180]


def _shape_texts(shape):
    try:
        if getattr(shape, "has_table", False):
            rows = []
            for row in shape.table.rows:
                cells = [(getattr(cell, "text", None) or "").strip() for cell in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            return rows
    except Exception:
        pass

    try:
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            texts = []
            for child in shape.shapes:
                texts.extend(_shape_texts(child))
            return texts
    except Exception:
        pass

    try:
        if getattr(shape, "has_text_frame", False):
            text = (shape.text_frame.text or "").strip()
            if text:
                return [text]
    except Exception:
        pass
    return []


def extract_pptx_text(data):
    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise CompareError(
            "Could not read the PowerPoint. Upload a .pptx file."
        ) from exc

    slides = []
    for index, slide in enumerate(presentation.slides, 1):
        bits = []
        try:
            shapes = list(slide.shapes)
        except Exception:
            shapes = []
        for shape in shapes:
            bits.extend(_shape_texts(shape))

        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
                if notes:
                    bits.append("Speaker notes: " + notes)
        except Exception:
            pass

        body = "\n".join(bit for bit in bits if bit)
        if body:
            slides.append(f"--- Slide {index} ---\n{body}")

    return "\n\n".join(slides).strip()


def _table_text(table):
    rows = []
    for row in table.rows:
        try:
            cells = [" ".join((cell.text or "").split()) for cell in row.cells]
        except Exception:
            continue
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def extract_docx_text(data):
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise CompareError(
            "Could not read a Word document. Upload a .docx file."
        ) from exc

    parts = []
    for child in document.element.body.iterchildren():
        try:
            if child.tag == qn("w:p"):
                text = Paragraph(child, document).text.strip()
                if text:
                    parts.append(text)
            elif child.tag == qn("w:tbl"):
                table_text = _table_text(Table(child, document))
                if table_text:
                    parts.append(table_text)
        except Exception:
            continue

    seen_hf = set()
    extra = []
    try:
        for section in document.sections:
            for hf in (section.header, section.footer):
                if hf is None:
                    continue
                for paragraph in hf.paragraphs:
                    text = paragraph.text.strip()
                    if text and text not in seen_hf:
                        seen_hf.add(text)
                        extra.append(text)
    except Exception:
        pass
    if extra:
        parts.append("Header/footer: " + " | ".join(extra))

    return "\n".join(parts).strip()


def _build_prompt(firm_config):
    ctx = (firm_config or {}).get("firm_context") or ""
    ctx = f" {ctx.strip()}" if ctx.strip() else ""
    return COMPARE_PROMPT.format(firm_context=ctx)


def _build_user_content(diagram_text, drafts):
    chunks = ["POWERPOINT DIAGRAM\n", diagram_text, "\n\nWORD DRAFTS\n"]
    for name, text in drafts:
        body = text or "(No extractable text.)"
        chunks.append(f"\n--- {safe_filename(name)} ---\n{body}\n")
    return "".join(chunks)


def payload_char_count(diagram_text, drafts):
    return len(_build_user_content(diagram_text, drafts))


def _clean_paragraph(raw):
    text = (raw or "").strip()
    if text.startswith("```"):
        text = text.split("\n", 1)[-1].rsplit("```", 1)[0].strip()
    return text


def compare_diagram_to_drafts(
    diagram_text,
    drafts,
    client,
    *,
    model=None,
    firm_id=None,
    firm_config=None,
):
    """Ask the model whether the Word drafts match the PowerPoint diagram.

    drafts is a list of (filename, extracted_text).
    Returns a single comparison paragraph.
    """
    diagram_text = (diagram_text or "").strip()
    if not diagram_text:
        raise CompareError(
            "No readable text was found in the PowerPoint. "
            "If the diagram is image-only, this version cannot compare it."
        )

    cleaned = []
    for name, text in drafts:
        cleaned.append((safe_filename(name), (text or "").strip()))
    if not cleaned:
        raise CompareError("Upload at least one Word document.")
    if not any(text for _, text in cleaned):
        raise CompareError(
            "No readable text was found in the Word document(s)."
        )

    user_content = _build_user_content(diagram_text, cleaned)
    log.info(
        "Comparing diagram (%d chars) against %d draft(s); payload %d chars",
        len(diagram_text),
        len(cleaned),
        len(user_content),
    )
    if len(user_content) > MAX_TOTAL_CHARS:
        raise CompareError(
            "These documents are too large to compare at once. "
            "Remove some Word documents and try again."
        )

    model = model or os.environ.get("OPENAI_MODEL", "gpt-5.6-terra")
    from ai_logger import log_ai_call, extract_xai_usage, completion_details

    call_start = time.time()
    raw = ""
    details = None
    try:
        response = client.chat.completions.create(
            model=model,
            messages=[
                {"role": "system", "content": _build_prompt(firm_config)},
                {"role": "user", "content": user_content},
            ],
        )
        raw = _clean_paragraph(response.choices[0].message.content)
        details = completion_details(response, raw)
        if not raw:
            log_ai_call(
                provider="openai",
                model=model,
                tool=TOOL_KEY,
                status="error",
                execution_ms=int((time.time() - call_start) * 1000),
                notes="Empty comparison text",
                firm_id=firm_id,
                **extract_xai_usage(response),
            )
            raise CompareError("The comparison came back empty. Please try again.")
        log_ai_call(
            provider="openai",
            model=model,
            tool=TOOL_KEY,
            status="success",
            execution_ms=int((time.time() - call_start) * 1000),
            firm_id=firm_id,
            **extract_xai_usage(response),
        )
        return raw
    except CompareError:
        raise
    except Exception as exc:
        log_ai_call(
            provider="openai",
            model=model,
            tool=TOOL_KEY,
            status="error",
            execution_ms=int((time.time() - call_start) * 1000),
            notes=traceback.format_exc(),
            firm_id=firm_id,
        )
        if details:
            exc.details = details
        raise
